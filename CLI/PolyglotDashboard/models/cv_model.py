import os
import sys

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class CurriculoModel:
    """Model de dados e gerador de currículos para o ecossistema Pedro Victor (UFF / ONS)."""

    def __init__(self):
        self.dados = {
            "nome": "Pedro Victor Rodrigues Veras",
            "localidade": "Niterói, Rio de Janeiro - RJ",
            "email": "pedrovictor.rveras12@gmail.com",
            "github": "https://github.com/PedroVic12",
            "linkedin": "https://www.linkedin.com/in/pedrovictor12/",
            "resumo": "Estudante de Engenharia Elétrica na Universidade Federal Fluminense (UFF), com forte atuação em desenvolvimento de software, automação de processos, desenvolvimento desktop Qt/QML e análise de sistemas de potência (ONS).",
            "formacao": [
                {"periodo": "2022 - Atual", "instituicao": "Universidade Federal Fluminense (UFF)", "curso": "Engenharia Elétrica"}
            ],
            "experiencia": [
                {"periodo": "2025 - 2026", "empresa": "Operador Nacional do Sistema Elétrico (ONS)", "cargo": "Estagiário PLC", "descricao": "Atuação no Planejamento da Operação de Curto Prazo do SIN, desenvolvimento de ferramentas internas de análise de rede e automação de relatórios."}
            ]
        }

    def gerar_pdf(self, vaga: str, objetivo: str, file_path: str) -> tuple[bool, str]:
        """Gera o arquivo PDF utilizando ReportLab com layout profissional."""
        if not REPORTLAB_AVAILABLE:
            return False, "Biblioteca 'reportlab' não está instalada no ambiente Python."

        try:
            doc = SimpleDocTemplate(
                file_path,
                pagesize=A4,
                rightMargin=40,
                leftMargin=40,
                topMargin=40,
                bottomMargin=40
            )
            styles = getSampleStyleSheet()

            # Estilos Customizados
            style_nome = ParagraphStyle(
                'Nome',
                parent=styles['Normal'],
                fontSize=22,
                leading=26,
                spaceAfter=6,
                textColor=colors.HexColor("#1e1e2e"),
                alignment=1,
                fontName='Helvetica-Bold'
            )
            style_contato = ParagraphStyle(
                'Contato',
                parent=styles['Normal'],
                fontSize=10,
                textColor=colors.HexColor("#585b70"),
                alignment=1,
                spaceAfter=15
            )
            style_secao = ParagraphStyle(
                'Secao',
                parent=styles['Normal'],
                fontSize=12,
                leading=14,
                textColor=colors.HexColor("#1e66f5"),
                spaceBefore=14,
                spaceAfter=6,
                fontName='Helvetica-Bold'
            )
            style_corpo = ParagraphStyle(
                'Corpo',
                parent=styles['Normal'],
                fontSize=10,
                textColor=colors.HexColor("#313244"),
                spaceAfter=8,
                leading=14
            )

            story = []

            # Cabeçalho
            story.append(Paragraph(self.dados['nome'], style_nome))
            contato_txt = f"{self.dados['localidade']} | {self.dados['email']}"
            story.append(Paragraph(contato_txt, style_contato))

            links_txt = f"GitHub: {self.dados['github']} | LinkedIn: {self.dados['linkedin']}"
            story.append(Paragraph(links_txt, style_contato))

            # Objetivo Profissional
            story.append(Paragraph("OBJETIVO PROFISSIONAL", style_secao))
            texto_obj = f"<b>Vaga Alvo: {vaga}</b><br/>{objetivo}" if vaga else objetivo
            story.append(Paragraph(texto_obj, style_corpo))

            # Resumo
            story.append(Paragraph("RESUMO PROFISSIONAL", style_secao))
            story.append(Paragraph(self.dados['resumo'], style_corpo))

            # Formação
            story.append(Paragraph("FORMAÇÃO ACADÊMICA", style_secao))
            for f in self.dados['formacao']:
                txt_f = f"<b>{f['curso']}</b> - {f['instituicao']} ({f['periodo']})"
                story.append(Paragraph(txt_f, style_corpo))

            # Experiência
            story.append(Paragraph("EXPERIÊNCIA PROFISSIONAL", style_secao))
            for e in self.dados['experiencia']:
                txt_e = f"<b>{e['cargo']}</b> - {e['empresa']} ({e['periodo']})<br/>{e['descricao']}"
                story.append(Paragraph(txt_e, style_corpo))

            doc.build(story)
            return True, f"PDF gerado com sucesso em: {file_path}"
        except Exception as err:
            return False, f"Erro ao gerar PDF: {str(err)}"

    def gerar_docx(self, vaga: str, objetivo: str, file_path: str) -> tuple[bool, str]:
        """Gera o arquivo Word (.docx) utilizando python-docx."""
        if not DOCX_AVAILABLE:
            return False, "Biblioteca 'python-docx' não está instalada no ambiente Python."

        try:
            doc = Document()
            h1 = doc.add_heading(self.dados['nome'], level=1)
            h1.alignment = 1

            p_contato = doc.add_paragraph(f"{self.dados['localidade']} | {self.dados['email']}")
            p_contato.alignment = 1

            doc.add_heading('OBJETIVO PROFISSIONAL', level=2)
            doc.add_paragraph(f"Vaga Alvo: {vaga}\n{objetivo}" if vaga else objetivo)

            doc.add_heading('RESUMO PROFISSIONAL', level=2)
            doc.add_paragraph(self.dados['resumo'])

            doc.save(file_path)
            return True, f"Word gerado com sucesso em: {file_path}"
        except Exception as err:
            return False, f"Erro ao gerar DOCX: {str(err)}"
